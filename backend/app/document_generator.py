import io
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .math_text import normalize_math_text


FONT_NAME = "BIZ UDPGothic"


def _set_run_font(run, *, size: float = 10.5, bold: bool = False):
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run.font.size = Pt(size)
    run.bold = bold


def _set_cell_margins(cell, top=60, start=0, bottom=60, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def _add_bottom_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    p_pr.append(borders)


def _add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, value, end])
    _set_run_font(run, size=8.5)


def _add_math_runs(paragraph, text: str):
    text = normalize_math_text(text)
    pattern = re.compile(r"(\^(?:\{([^{}]+)\}|\(([^()]+)\)|([0-9]+))|_(?:\{([^{}]+)\}|([0-9]+)))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            _set_run_font(run)
        superscript = match.group(2) or match.group(3) or match.group(4)
        subscript = match.group(5) or match.group(6)
        run = paragraph.add_run(superscript or subscript)
        _set_run_font(run, size=8)
        run.font.superscript = superscript is not None
        run.font.subscript = subscript is not None
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        _set_run_font(run)


def _distribute(total: int, count: int) -> list[int]:
    quotient, remainder = divmod(total, count)
    return [quotient + (1 if index < remainder else 0) for index in range(count)]


def build_test_docx(*, school_name: str, grade: str, subject: str, title: str,
                    duration_minutes: int, total_points: int, problems: list[dict]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(16)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(10.5)

    school = document.add_paragraph()
    school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    school.paragraph_format.space_after = Pt(2)
    _set_run_font(school.add_run(school_name), size=10.5)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(8)
    _set_run_font(heading.add_run(f"{grade}　{subject}\n{title}"), size=15, bold=True)

    info = document.add_table(rows=2, cols=3)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    widths = [Mm(52), Mm(80), Mm(42)]
    values = [
        ["実施日：____年__月__日", f"制限時間：{duration_minutes}分", f"満点：{total_points}点"],
        [f"{grade.replace('高', '')}年 ____組　____番", "氏名 ________________________________", f"得点 ______ / {total_points}"],
    ]
    for row_index, row in enumerate(info.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = widths[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if column_index == 2 else WD_ALIGN_PARAGRAPH.LEFT
            _set_run_font(paragraph.add_run(values[row_index][column_index]), size=9.5)
    _remove_table_borders(info)

    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(9)
    _add_bottom_rule(rule)

    groups = [problems[index:index + 4] for index in range(0, len(problems), 4)]
    points = _distribute(total_points, len(groups))
    for major_index, group in enumerate(groups, 1):
        major = document.add_paragraph()
        major.paragraph_format.page_break_before = major_index > 1
        major.paragraph_format.keep_with_next = True
        major.paragraph_format.space_before = Pt(4)
        major.paragraph_format.space_after = Pt(5)
        _set_run_font(major.add_run(f"第{major_index}問　次の問いに答えなさい。　　　　　　　　　〔{points[major_index - 1]}点〕"), size=11, bold=True)
        for sub_index, problem in enumerate(group, 1):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.line_spacing = 1.25
            paragraph.paragraph_format.space_after = Pt(42)
            prefix = paragraph.add_run(f"({sub_index})　")
            _set_run_font(prefix)
            _add_math_runs(paragraph, problem["body"].strip())

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_field(footer, "PAGE")
    _set_run_font(footer.add_run(" / "), size=8.5)
    _add_field(footer, "NUMPAGES")

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
