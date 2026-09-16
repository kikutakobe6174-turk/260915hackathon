"""Gemini解析で得た実配点が、PDF・Word・プレビュー用APIで一致することを検証する。"""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.main import app

from test_llm_flow import FakeProvider


def _generate_and_save(client):
    import base64

    encoded = base64.b64encode(b"test image bytes").decode()
    analysis = client.post("/llm/trend-draft", json={
        "test_id": 1, "user_id": 1, "image_base64": encoded, "media_type": "image/jpeg",
    })
    assert analysis.status_code == 200, analysis.text
    analysis_job_id = analysis.json()["job_id"]
    assert client.post(f"/llm/trend-drafts/{analysis_job_id}/confirm", json={"user_id": 1}).status_code == 200
    generated = client.post(f"/llm/trend-drafts/{analysis_job_id}/generate", json={"user_id": 1})
    assert generated.status_code == 200, generated.text
    return generated.json()["job_id"], generated.json()["problems"]


def test_real_points_are_used_everywhere(tmp_path, monkeypatch):
    # FakeProvider の解析結果は 5点 + 10点 + 5点 = 20点。等分なら 7/7/6 になるはず。
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "points.db"))
    import app.main as main

    app.dependency_overrides[main.get_provider] = lambda: FakeProvider()
    try:
        with TestClient(app) as client:
            generation_job_id, drafts = _generate_and_save(client)

            # 1) 生成直後の下書きからのPDF出力（満点100を指定しても実配点20が優先される）
            pdf = client.post(f"/llm/problem-batches/{generation_job_id}/pdf", json={
                "duration_minutes": 50, "total_points": 100,
            })
            assert pdf.status_code == 200, pdf.text
            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf.content)).pages)
            assert "満点：20点" in pdf_text
            assert "得点 ______ / 20" in pdf_text
            assert "〔20点〕" in pdf_text  # 大問は小問配点の合計
            assert "（5点）" in pdf_text and "（10点）" in pdf_text

            # 2) 同じ内容のWord出力
            word = client.post(f"/llm/problem-batches/{generation_job_id}/word", json={
                "duration_minutes": 50, "total_points": 100,
            })
            assert word.status_code == 200, word.text
            document = Document(BytesIO(word.content))
            word_text = "\n".join(p.text for p in document.paragraphs)
            table_text = "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
            assert "〔20点〕" in word_text
            assert "（5点）" in word_text and "（10点）" in word_text
            assert "満点：20点" in table_text

            # 3) 問題バンクへ保存したあとの、テスト単位のPDF/Word出力
            saved = client.post(f"/llm/problem-batches/{generation_job_id}/save", json={"user_id": 1, "problems": drafts})
            assert saved.status_code == 200, saved.text
            problem_ids = saved.json()["problem_ids"]

            points = client.get("/tests/1/problem-points").json()
            assert points["source"] == "analysis"
            assert points["total_points"] == 20
            assert sorted(item["points"] for item in points["items"]) == [5, 5, 10]
            assert {item["problem_id"] for item in points["items"]} == set(problem_ids)

            test_pdf = client.post("/tests/1/pdf", json={
                "problem_ids": problem_ids, "duration_minutes": 50, "total_points": 100,
            })
            assert test_pdf.status_code == 200, test_pdf.text
            test_pdf_text = "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(test_pdf.content)).pages)
            assert "満点：20点" in test_pdf_text
            assert "（10点）" in test_pdf_text

            test_word = client.post("/tests/1/word", json={
                "problem_ids": problem_ids, "duration_minutes": 50, "total_points": 100,
            })
            assert test_word.status_code == 200, test_word.text
            assert "（10点）" in "\n".join(p.text for p in Document(BytesIO(test_word.content)).paragraphs)

            # 4) プレビューが使う配点は、PDFの大問配点と同じ合計になる
            assert sum(item["points"] for item in points["items"]) == 20
    finally:
        app.dependency_overrides.clear()


def test_falls_back_to_even_distribution_without_analysis(tmp_path, monkeypatch):
    """解析配点が無い手入力の問題では、従来どおり満点を大問へ等分する（既存挙動を壊さない）。"""
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "even.db"))
    with TestClient(app) as client:
        problem_ids = []
        for index in range(4):
            response = client.post("/problems", json={
                "unit_id": 1, "format_id": 1, "difficulty": 1,
                "body": f"{index + 1}x + 1 を計算しなさい。", "answer": "解答", "explanation": "解説",
                "is_return": False, "hints": ["a", "b", "c"], "prerequisite_unit_ids": [], "source": "manual",
            })
            assert response.status_code == 200, response.text
            problem_ids.append(response.json()["id"])

        points = client.get("/tests/1/problem-points").json()
        assert points["source"] == "even"
        assert points["total_points"] is None

        pdf = client.post("/tests/1/pdf", json={
            "problem_ids": problem_ids, "duration_minutes": 50, "total_points": 100,
        })
        assert pdf.status_code == 200, pdf.text
        text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf.content)).pages)
        assert "満点：100点" in text
        assert "〔100点〕" in text
        assert "点）" not in text  # 小問ごとの配点は表示しない


def test_saved_points_survive_reanalysis(tmp_path, monkeypatch):
    """問題バンク保存時に配点を確定させるので、同じテストを解析し直しても出力の配点はぶれない。"""
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "reanalysis.db"))
    import app.main as main

    app.dependency_overrides[main.get_provider] = lambda: FakeProvider()
    try:
        with TestClient(app) as client:
            generation_job_id, drafts = _generate_and_save(client)
            saved = client.post(f"/llm/problem-batches/{generation_job_id}/save", json={"user_id": 1, "problems": drafts})
            assert saved.status_code == 200, saved.text
            before = client.get("/tests/1/problem-points").json()
            assert before["total_points"] == 20

            # 出題傾向を別の配点で上書きする（＝テストを解析し直した状態）
            current = client.get("/tests/1/trends").json()["items"]
            rewritten = client.put("/tests/1/trends", json={
                "user_id": 1,
                "items": [{
                    "unit_id": item["unit_id"], "format_id": item["format_id"],
                    "question_no": item["question_no"], "points": 40,
                    "difficulty": item["difficulty"], "confidence": item["confidence"],
                    "source": "manual", "llm_job_id": None,
                } for item in current],
            })
            assert rewritten.status_code == 200, rewritten.text

            after = client.get("/tests/1/problem-points").json()
            assert after["source"] == "analysis"
            assert after["total_points"] == before["total_points"] == 20
            assert after["items"] == before["items"]
    finally:
        app.dependency_overrides.clear()
